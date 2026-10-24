"""Write reviewer decisions from the mapping-review workbook back to the vocabulary.

Each decided row on the workbook's "Review" sheet is matched to its
keywords_v1.csv row by (keyword, mapped concept_id) among imported rows:

- accept: keep the mapping and set review_status.
- change: map the keyword to "Corrected concept_id" and set review_status.
- move to policy: remove it from the vocabulary and append it to
  policy_candidates_ldnoobw_en.csv with its "Policy route".
- drop: remove it from the vocabulary.

"Set context_required" (true/false) optionally overrides the flag on kept rows.
Rows without a decision are left untouched. Every decision is validated before
any file is written, so an invalid workbook changes nothing.

Install: pip install openpyxl
Run: python apply_review.py review/ldnoobw_mapping_review.xlsx --status human_reviewed
"""

import argparse
import csv

import build
from matching import Matcher, norm

POLICY_FILE = "policy_candidates_ldnoobw_en.csv"
POLICY_ROUTES = (
    "minor_safety",
    "non_consent",
    "incest",
    "bestiality",
    "violence_extremism",
    "solicitation",
    "hate_slur",
)
DECISIONS = ("accept", "change", "move to policy", "drop")
COLUMNS = {
    "keyword": "Keyword",
    "concept": "Mapped concept_id",
    "decision": "Decision",
    "corrected": "Corrected concept_id",
    "route": "Policy route",
    "context": "Set context_required",
}


def read_decisions(workbook_path):
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows = workbook["Review"].iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
        missing = [name for name in COLUMNS.values() if name not in header]
        if missing:
            raise ValueError(f"Review sheet is missing columns: {missing}")
        index = {key: header.index(name) for key, name in COLUMNS.items()}
        decided = []
        for row in rows:
            record = {
                key: "" if i >= len(row) or row[i] is None else str(row[i]).strip()
                for key, i in index.items()
            }
            if record["decision"]:
                decided.append(record)
        return decided
    finally:
        workbook.close()


def read_csv(path):
    with path.open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        return reader.fieldnames, list(reader)


def apply_review(workbook_path, status, dry_run=False):
    """Validate every decision, then rewrite the vocabulary and policy files."""
    taxonomy = build.load_taxonomy()
    nodes = {node["id"]: node for node in taxonomy["nodes"]}
    vocab_path = build.ROOT / build.VOCABULARY_FILE
    policy_path = build.ROOT / POLICY_FILE
    vocab_fields, vocabulary = read_csv(vocab_path)
    policy_fields, policy = read_csv(policy_path)
    imported = {
        (norm(row["keyword"]), row["concept_id"]): row
        for row in vocabulary
        if row["source"] != build.SEED_SOURCE
    }

    removed, added_policy = set(), []
    summary = {decision: 0 for decision in DECISIONS}
    for d in read_decisions(workbook_path):
        key = (norm(d["keyword"]), d["concept"])
        where = f"{d['keyword']!r} ({d['concept']})"
        if d["decision"] not in DECISIONS:
            raise ValueError(f"Unknown decision {d['decision']!r} for {where}")
        if key not in imported:
            raise ValueError(f"No imported vocabulary row for {where}")
        if d["context"] not in ("", "true", "false"):
            raise ValueError(f"Set context_required must be true or false for {where}")
        row = imported[key]
        if d["decision"] == "change":
            target = nodes.get(d["corrected"])
            if target is None or target["level"] == 0:
                raise ValueError(
                    f"Corrected concept_id must be an L1-L3 concept for {where}"
                )
            if target["category_id"] == "IDENTITY":
                raise ValueError(f"Keywords may not map to IDENTITY concepts: {where}")
            row["concept_id"] = d["corrected"]
        if d["decision"] == "move to policy":
            if d["route"] not in POLICY_ROUTES:
                raise ValueError(
                    f"Policy route must be one of {POLICY_ROUTES} for {where}"
                )
            added_policy.append(
                {
                    "keyword": row["keyword"],
                    "language": row["language"],
                    "policy_route": d["route"],
                    "source": row["source"],
                    "review_status": "policy_review_required",
                }
            )
        if d["decision"] in ("move to policy", "drop"):
            removed.add(id(row))
        else:
            row["review_status"] = status
            if d["context"]:
                row["context_required"] = d["context"]
        summary[d["decision"]] += 1

    kept = [row for row in vocabulary if id(row) not in removed]
    build.check_vocabulary(kept)
    parents = {node_id: node["parent_id"] for node_id, node in nodes.items()}
    Matcher(parents, kept)
    known_policy = {(norm(r["keyword"]), r["language"]) for r in policy}
    new_policy = [
        r
        for r in added_policy
        if (norm(r["keyword"]), r["language"]) not in known_policy
    ]
    print(
        ("Would apply" if dry_run else "Applied"),
        ", ".join(f"{count} {decision}" for decision, count in summary.items()),
    )
    if dry_run:
        return summary
    # Same dialects as the generator: CRLF vocabulary, LF policy candidates.
    with vocab_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=vocab_fields)
        writer.writeheader()
        writer.writerows(kept)
    with policy_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=policy_fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(policy + new_policy)
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook")
    parser.add_argument(
        "--status",
        default="human_reviewed",
        help="review_status for accepted and changed rows",
    )
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    apply_review(args.workbook, args.status, args.dry_run)
